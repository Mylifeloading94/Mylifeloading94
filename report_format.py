"""
The house spreadsheet format for every backtest deliverable in this repo.

The owner supplied a reference workbook
(`mylifeloading_portfolio_gold_fx_100k_backtest.xlsx`) and asked that every
spreadsheet follow it. This module is that format, implemented once, so a
deliverable is a call to `build_workbook()` rather than a fresh guess at a
layout. `.claude/commands/spreadsheet-format.md` documents the contract.

Three sheets, always, in this order:

  Summary    title / subtitle / a prose METHODOLOGY block that carries the
             caveats, then the headline metric row, then Concurrency, then
             the Per-Instrument Breakdown.
  Trade Log  one row per trade, 18 fixed columns, Date first, running
             Balance ($) last before the Week/Month grouping keys.
  Periods    Daily then Weekly then Monthly, each a
             Period / P&L ($) / % of Starting Balance block.

The prose block is not decoration. The reference workbook put its
overfitting caveat, its concurrency risk and its lot-rounding disclosure
there, above the numbers, and that is the point of it: the reader meets the
caveats before the ROI.
"""
from __future__ import annotations

import datetime as _dt

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- styling ---
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(size=10, italic=True)
SECTION_FONT = Font(bold=True, size=11)
HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F4F5F")
PROSE_FONT = Font(size=9)
WIN_FILL = PatternFill("solid", fgColor="E8F5E9")
LOSS_FILL = PatternFill("solid", fgColor="FDECEA")
THIN = Side(style="thin", color="D0D0D0")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '#,##0.00;[Red]-#,##0.00'
PCT = '0.00%'
PRICE = '0.#####'

TRADE_LOG_COLUMNS = [
    ("Date", 12, None),
    ("Entry Time (UTC)", 18, "yyyy-mm-dd hh:mm"),
    ("Pair", 10, None),
    ("Direction", 10, None),
    ("Entry", 12, PRICE),
    ("Stop", 12, PRICE),
    ("Target", 12, PRICE),
    ("Exit", 12, PRICE),
    ("Outcome", 12, None),
    ("Stop (pips)", 11, '0.0'),
    ("Result (pips)", 12, '0.0;[Red]-0.0'),
    ("Result R", 10, '0.000;[Red]-0.000'),
    ("Win", 6, '0'),
    ("Lots", 9, '0.00'),
    ("P&L ($)", 12, MONEY),
    ("Balance ($)", 14, '#,##0.00'),
    ("Week", 10, None),
    ("Month", 10, None),
]


def _write_header_row(ws, row: int, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CELL_BORDER
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def _prose(ws, row: int, text: str, width_cols: int = 10, height: int | None = None):
    """The methodology / caveat block. One merged cell, wrapped, above the numbers."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width_cols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = PROSE_FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = height or min(400, 14 * (text.count("\n") + text.count(". ") // 2 + 6))
    return row + 1


# ----------------------------------------------------------------- sheets ---
def build_summary(
    ws,
    *,
    title: str,
    subtitle: str,
    methodology: str,
    headline: dict,
    concurrency: dict | None = None,
    per_instrument: pd.DataFrame | None = None,
):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=2, column=1, value=subtitle).font = SUBTITLE_FONT
    row = _prose(ws, 4, methodology)
    row += 1

    labels = list(headline.keys())
    _write_header_row(ws, row, labels)
    for i, key in enumerate(labels, start=1):
        c = ws.cell(row=row + 1, column=i, value=headline[key])
        c.border = CELL_BORDER
        low = key.lower()
        if "rate" in low or "roi" in low or "drawdown" in low or low.endswith("%"):
            c.number_format = PCT
        elif "balance" in low or "p&l" in low or "profit ($)" in low:
            c.number_format = '#,##0.00'
        elif "factor" in low:
            c.number_format = '0.000'
    row += 3

    if concurrency:
        ws.cell(row=row, column=1, value="Concurrency").font = SECTION_FONT
        row += 1
        for k, v in concurrency.items():
            ws.cell(row=row, column=1, value=k)
            ws.cell(row=row, column=2, value=v)
            row += 1
        row += 1

    if per_instrument is not None and len(per_instrument):
        ws.cell(row=row, column=1, value="Per-Instrument Breakdown").font = SECTION_FONT
        row += 1
        cols = list(per_instrument.columns)
        _write_header_row(ws, row, cols)
        for j, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(j)].width = 34 if "Param" in col else 14
        row += 1
        for _, r in per_instrument.iterrows():
            for j, col in enumerate(cols, start=1):
                c = ws.cell(row=row, column=j, value=r[col])
                c.border = CELL_BORDER
                low = col.lower()
                if "win rate" in low:
                    c.number_format = PCT
                elif "p&l" in low or "pnl" in low:
                    c.number_format = MONEY
                elif low == "pf":
                    c.number_format = '0.00'
            row += 1
    ws.freeze_panes = "A1"


def build_trade_log(ws, trades: pd.DataFrame, *, title: str):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    labels = [c[0] for c in TRADE_LOG_COLUMNS]
    widths = [c[1] for c in TRADE_LOG_COLUMNS]
    _write_header_row(ws, 3, labels, widths)

    for i, (_, t) in enumerate(trades.iterrows()):
        r = 4 + i
        for j, (name, _w, fmt) in enumerate(TRADE_LOG_COLUMNS, start=1):
            c = ws.cell(row=r, column=j, value=t.get(name))
            c.border = CELL_BORDER
            if fmt:
                c.number_format = fmt
        won = bool(t.get("Win"))
        fill = WIN_FILL if won else LOSS_FILL
        for j in (9, 11, 12, 13, 15):
            ws.cell(row=r, column=j).fill = fill
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:R{3 + len(trades)}"


def build_periods(ws, *, daily: pd.DataFrame, weekly: pd.DataFrame, monthly: pd.DataFrame,
                  title: str = "Daily / Weekly / Monthly Gains -- Portfolio Total"):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22

    row = 3
    for label, frame in (("Daily Gains", daily), ("Weekly Gains", weekly), ("Monthly Gains", monthly)):
        ws.cell(row=row, column=1, value=label).font = SECTION_FONT
        row += 1
        _write_header_row(ws, row, ["Period", "P&L ($)", "% of Starting Balance"])
        row += 1
        for _, r in frame.iterrows():
            ws.cell(row=row, column=1, value=r["Period"]).border = CELL_BORDER
            c = ws.cell(row=row, column=2, value=float(r["P&L ($)"]))
            c.number_format = MONEY
            c.border = CELL_BORDER
            c = ws.cell(row=row, column=3, value=float(r["% of Starting Balance"]))
            c.number_format = PCT
            c.border = CELL_BORDER
            row += 1
        row += 2
    ws.freeze_panes = "A2"


# ------------------------------------------------------------- assembling ---
def periods_from_trades(trades: pd.DataFrame, starting_balance: float):
    """Daily/Weekly/Monthly P&L as a fraction of the STARTING balance.

    The reference workbook expresses every period as a share of the opening
    balance, not of that period's own opening balance, so the three tables
    stay on one comparable scale and sum to the total return.
    """
    df = trades.copy()
    df["Date"] = pd.to_datetime(df["Date"])
    out = []
    for keyer in (
        lambda d: d.dt.strftime("%Y-%m-%d"),
        lambda d: d.dt.strftime("%G-W%V"),
        lambda d: d.dt.strftime("%Y-%m"),
    ):
        g = df.groupby(keyer(df["Date"]))["P&L ($)"].sum().reset_index()
        g.columns = ["Period", "P&L ($)"]
        g["% of Starting Balance"] = g["P&L ($)"] / starting_balance
        out.append(g)
    return out


def build_workbook(
    path: str,
    *,
    title: str,
    subtitle: str,
    methodology: str,
    headline: dict,
    trades: pd.DataFrame,
    starting_balance: float,
    concurrency: dict | None = None,
    per_instrument: pd.DataFrame | None = None,
    trade_log_title: str | None = None,
):
    """Write the house-format workbook. `trades` must carry TRADE_LOG_COLUMNS names."""
    daily, weekly, monthly = periods_from_trades(trades, starting_balance)

    wb = Workbook()
    build_summary(
        wb.active, title=title, subtitle=subtitle, methodology=methodology,
        headline=headline, concurrency=concurrency, per_instrument=per_instrument,
    )
    wb.active.title = "Summary"
    build_trade_log(wb.create_sheet("Trade Log"), trades,
                    title=trade_log_title or f"Combined Trade Log -- {subtitle}")
    build_periods(wb.create_sheet("Periods"), daily=daily, weekly=weekly, monthly=monthly)
    wb.save(path)
    return path
