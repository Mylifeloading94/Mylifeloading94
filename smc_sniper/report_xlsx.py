"""Excel report builder (openpyxl).

Sheets: Summary, Verdict, Trade Ledger, Per-Pair, Per-Session, Monthly, Weekly,
In/Out-of-sample split, Walk-forward, Rejected setups, Matched-R control,
Perturbation, Data coverage.

The Verdict sheet is written in plain English and is required to be blunt about
whether the 68% win-rate target was reached.
"""
from __future__ import annotations

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_frame(ws, frame: pd.DataFrame, start_row: int = 1, title: str = "") -> int:
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = TITLE_FONT
        row += 2
    if frame is None or frame.empty:
        ws.cell(row=row, column=1, value="(no rows)")
        return row + 2
    for j, col in enumerate(frame.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(col))
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BORDER
    for i, (_, data) in enumerate(frame.iterrows(), start=row + 1):
        for j, col in enumerate(frame.columns, start=1):
            val = data[col]
            if isinstance(val, (pd.Timestamp,)):
                val = str(val)
            elif hasattr(val, "item"):
                try:
                    val = val.item()
                except (ValueError, AttributeError):
                    val = str(val)
            elif not isinstance(val, (int, float, str, type(None))):
                val = str(val)
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = BORDER
            if isinstance(val, float):
                cell.number_format = "0.0000" if abs(val) < 10 else "#,##0.00"
    for j, col in enumerate(frame.columns, start=1):
        width = max(len(str(col)) + 2,
                    *(len(str(v)[:28]) + 2 for v in frame[col].head(200)))
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 9), 30)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return row + len(frame) + 3


def _write_lines(ws, lines: list, start_row: int = 1, title: str = "") -> int:
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = TITLE_FONT
        row += 2
    for line in lines:
        cell = ws.cell(row=row, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if line and (line.isupper() or line.endswith(":")):
            cell.font = BOLD
        row += 1
    ws.column_dimensions["A"].width = 118
    return row + 1


def build(path: str, *, summary: pd.DataFrame, verdict: list[str],
          ledger: pd.DataFrame, per_pair: pd.DataFrame, per_session: pd.DataFrame,
          monthly: pd.DataFrame, weekly: pd.DataFrame, splits: pd.DataFrame,
          walk_forward: pd.DataFrame, rejected: pd.DataFrame,
          matched_r: pd.DataFrame | None = None,
          perturbation: pd.DataFrame | None = None,
          coverage: pd.DataFrame | None = None,
          per_setup: pd.DataFrame | None = None,
          adaptive: dict | None = None,
          stack_comparison: pd.DataFrame | None = None) -> str:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Summary")
    row = 1
    if stack_comparison is not None and not stack_comparison.empty:
        row = _write_frame(ws, stack_comparison, row,
                           "Stack comparison -- both timeframe stacks, side by side")
    _write_frame(ws, summary, row, "SMC Sniper -- Backtest Summary")

    ws = wb.create_sheet("Verdict")
    _write_lines(ws, verdict, 1, "Verdict -- plain English")

    _write_frame(wb.create_sheet("Trade Ledger"), ledger, 1, "Every trade taken")
    _write_frame(wb.create_sheet("Per-Pair"), per_pair, 1, "Performance by pair")
    _write_frame(wb.create_sheet("Per-Session"), per_session, 1, "Performance by session")
    if per_setup is not None:
        _write_frame(wb.create_sheet("Per-Setup"), per_setup, 1, "Performance by setup type")
    _write_frame(wb.create_sheet("Monthly"), monthly, 1, "Monthly performance")
    _write_frame(wb.create_sheet("Weekly"), weekly, 1, "Weekly performance")
    _write_frame(wb.create_sheet("In-Out-Sample"), splits, 1,
                 "In-sample vs out-of-sample (chronological split)")
    _write_frame(wb.create_sheet("Walk-Forward"), walk_forward, 1,
                 "Walk-forward folds -- pairs selected on each fit window only")

    if matched_r is not None:
        ws = wb.create_sheet("Matched-R Control")
        row = _write_lines(ws, [
            "Anti-TP-shrinking control.",
            "Same entries, management stripped, one flat target at each R.",
            "If a win rate came from a small target rather than real edge, win",
            "rate collapses as R rises while expectancy stays flat or negative.",
            "'breakeven_wr_needed' is the win rate required to break even at that R.",
        ], 1, "Matched-R control")
        _write_frame(ws, matched_r, row)

    if perturbation is not None:
        ws = wb.create_sheet("Perturbation")
        row = _write_lines(ws, [
            "Robustness check: each parameter nudged, everything else held.",
            "A result that survives only at one exact value is a curve fit.",
        ], 1, "Parameter perturbation")
        _write_frame(ws, perturbation, row)

    _write_frame(wb.create_sheet("Rejected Setups"), rejected, 1,
                 "Why setups were declined (the selectivity audit)")

    if adaptive:
        ws = wb.create_sheet("Adaptive")
        row = 1
        for dim, table in adaptive.items():
            if isinstance(table, pd.DataFrame) and not table.empty:
                row = _write_frame(ws, table, row, f"By {dim}")

    if coverage is not None:
        _write_frame(wb.create_sheet("Data Coverage"), coverage, 1,
                     "Bars available per symbol and timeframe")

    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
